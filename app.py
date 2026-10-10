import os, re, io, base64, secrets, hmac, json, csv, unicodedata, zipfile
from datetime import datetime, timedelta, timezone
from functools import wraps
from pathlib import Path
from copy import deepcopy
from flask import Flask, render_template, request, session, redirect, url_for, jsonify, send_file, abort
from sqlalchemy import create_engine, MetaData, Table, Column, Integer, String, LargeBinary, DateTime, select, delete, func, text
from sqlalchemy.exc import IntegrityError
from PIL import Image, UnidentifiedImageError
from docx import Document
from docx.shared import Inches
from docx.oxml import OxmlElement
from docx.table import Table as WordTable
from flask import flash

ROOT = Path(__file__).parent
app = Flask(__name__)
for key in ('SECRET_KEY', 'ADMIN_PASSWORD', 'DELETE_PASSWORD'):
    if not os.environ.get(key):
        raise RuntimeError(f'Configura la variable {key} antes de iniciar.')
app.config.update(SECRET_KEY=os.environ['SECRET_KEY'], MAX_CONTENT_LENGTH=5_000_000,
                  SESSION_COOKIE_HTTPONLY=True, SESSION_COOKIE_SAMESITE='Lax',
                  SESSION_COOKIE_SECURE=bool(os.environ.get('RENDER')), PERMANENT_SESSION_LIFETIME=timedelta(hours=2))
url = os.environ.get('DATABASE_URL', 'sqlite:///' + str(ROOT / 'local.db'))
if os.environ.get('RENDER') and not os.environ.get('DATABASE_URL'):
    raise RuntimeError('En Render debes configurar DATABASE_URL para conservar los registros.')
if url.startswith('postgres://'): url = url.replace('postgres://', 'postgresql://', 1)
if url.startswith('postgresql://'):
    url = url.replace('postgresql://', 'postgresql+psycopg2://', 1)
elif url.startswith('postgresql+psycopg://'):
    url = url.replace('postgresql+psycopg://', 'postgresql+psycopg2://', 1)
engine = create_engine(url, pool_pre_ping=True)
meta = MetaData()
records = Table('firmantes', meta, Column('id', Integer, primary_key=True),
    Column('nombre', String(180), nullable=False), Column('rut', String(12), unique=True, nullable=False),
    Column('firma', LargeBinary, nullable=False), Column('fecha', DateTime(timezone=True), nullable=False))
members = Table('socios_habilitados', meta,
    Column('rut', String(12), primary_key=True), Column('nombre', String(180), nullable=False))
roster_init = Table('nomina_inicializada', meta, Column('id', Integer, primary_key=True))
with engine.begin() as conn:
    if engine.dialect.name == 'postgresql':
        conn.execute(text('SELECT pg_advisory_xact_lock(54002610)'))
    meta.create_all(conn)


def csrf():
    if 'csrf' not in session: session['csrf'] = secrets.token_hex(32)
    return session['csrf']
app.jinja_env.globals['csrf'] = csrf

@app.before_request
def protect():
    if request.method == 'POST':
        token = request.headers.get('X-CSRF-Token') or request.form.get('csrf', '')
        if not hmac.compare_digest(str(token), session.get('csrf', 'missing')): abort(403)

@app.after_request
def headers(response):
    response.headers['Cache-Control'] = 'no-store'
    response.headers['X-Content-Type-Options'] = 'nosniff'
    response.headers['X-Frame-Options'] = 'DENY'
    response.headers['Referrer-Policy'] = 'same-origin'
    response.headers['Content-Security-Policy'] = "default-src 'self'; img-src 'self' data:; style-src 'self'; script-src 'self'; frame-ancestors 'none'; base-uri 'self'; form-action 'self'"
    return response

def admin(fn):
    @wraps(fn)
    def wrapper(*a, **kw):
        if not session.get('admin'): return redirect(url_for('login'))
        return fn(*a, **kw)
    return wrapper

def normalize_rut(value):
    raw = re.sub(r'[.\s-]', '', value).upper()
    if not re.fullmatch(r'\d{7,8}[0-9K]', raw): raise ValueError('Ingresa un RUT válido.')
    body, dv = raw[:-1], raw[-1]
    total = sum(int(n) * (2 + i % 6) for i, n in enumerate(reversed(body)))
    result = 11 - total % 11
    expected = '0' if result == 11 else 'K' if result == 10 else str(result)
    if dv != expected or int(body) == 0: raise ValueError('El dígito verificador del RUT no es válido.')
    b = body.zfill(8)
    return f'{b[:2]}.{b[2:5]}.{b[5:]}-{dv}'

def name_key(value):
    text = unicodedata.normalize('NFKD', value).encode('ascii', 'ignore').decode().upper()
    return sorted(re.findall(r'[A-Z]+', text))

# Seed once, atomically. Subsequent administrator roster changes survive restarts.
roster_path = ROOT / 'nomina.json'
if roster_path.exists():
    initial = json.loads(roster_path.read_text(encoding='utf-8'))
    for item in initial:
        item['rut'] = normalize_rut(item['rut'])
    with engine.begin() as conn:
        if engine.dialect.name == 'postgresql':
            conn.execute(text('SELECT pg_advisory_xact_lock(54002610)'))
        if conn.scalar(select(roster_init.c.id).where(roster_init.c.id == 1)) is None:
            # Preserve an administrator-imported roster when upgrading.
            if conn.scalar(select(func.count()).select_from(members)) == 0:
                conn.execute(members.insert(), initial)
            conn.execute(roster_init.insert().values(id=1))

def signature(value):
    if not isinstance(value, str) or not value.startswith('data:image/png;base64,'): raise ValueError('Debes dibujar tu firma.')
    try:
        data = base64.b64decode(value.split(',', 1)[1], validate=True)
        im = Image.open(io.BytesIO(data))
        if im.format != 'PNG' or im.width > 2000 or im.height > 1000 or im.width < 100 or im.height < 50:
            raise ValueError('Firma de tamaño inválido.')
        rgba = im.convert('RGBA'); white = Image.new('RGBA', rgba.size, 'white'); white.alpha_composite(rgba)
        gray = white.convert('L')
        if sum(gray.histogram()[:200]) < 30: raise ValueError('Debes dibujar una firma visible.')
        out = io.BytesIO(); white.convert('RGB').save(out, format='PNG'); return out.getvalue()
    except (UnidentifiedImageError, OSError, base64.binascii.Error): raise ValueError('La imagen de firma no es válida.')

@app.get('/')
def index():
    document = Document(ROOT / 'plantillas/censura.docx')
    paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
    with engine.connect() as conn:
        authorized_members = conn.execute(select(members.c.nombre, members.c.rut).order_by(members.c.nombre, members.c.rut)).mappings().all()
    return render_template('index.html', document_paragraphs=paragraphs, authorized_members=authorized_members)

@app.post('/registrar')
def register():
    data = request.get_json(silent=True) or {}
    try:
        nombre = ' '.join(str(data.get('nombre', '')).split())
        if len(nombre) > 180 or len(nombre.split()) < 2: raise ValueError('Ingresa tu nombre completo y apellidos.')
        if data.get('consentimiento') is not True: raise ValueError('Confirma que deseas respaldar la solicitud.')
        rut = normalize_rut(str(data.get('rut', ''))); firma = signature(data.get('firma'))
        with engine.begin() as conn:
            member = conn.execute(select(members).where(members.c.rut == rut)).mappings().first()
            if member is None:
                return jsonify(message='El RUT no está habilitado en la nómina de socios. Contacta al administrador.'), 403
            if name_key(nombre) != name_key(member['nombre']):
                return jsonify(message='El nombre completo no coincide con la nómina para ese RUT. Revisa tus nombres y apellidos o contacta al administrador.'), 400
            conn.execute(records.insert().values(nombre=member['nombre'], rut=rut, firma=firma, fecha=datetime.now(timezone.utc)))
        return jsonify(message='Registro guardado. Gracias por firmar.'), 201
    except ValueError as e: return jsonify(message=str(e)), 400
    except IntegrityError: return jsonify(message='Este RUT ya tiene un registro. Contacta al administrador para corregirlo.'), 409

@app.route('/admin/login', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        now = datetime.now(timezone.utc).timestamp()
        if session.get('blocked_until', 0) > now:
            flash('Espera unos minutos antes de volver a intentar.'); return render_template('login.html'), 429
        if hmac.compare_digest(request.form.get('password',''), os.environ['ADMIN_PASSWORD']):
            session.clear(); session['admin'] = True; session.permanent = True
            return redirect(url_for('dashboard'))
        attempts = session.get('attempts', 0) + 1; session['attempts'] = attempts
        if attempts >= 5: session['blocked_until'] = now + 300; session['attempts'] = 0
        flash('Clave incorrecta.')
    return render_template('login.html')

@app.get('/admin')
@admin
def dashboard():
    page = max(1, request.args.get('page', 1, type=int)); q = request.args.get('q', '').strip()[:180]
    criteria = records.c.nombre.ilike('%'+q.replace('%','\\%').replace('_','\\_')+'%', escape='\\') | records.c.rut.contains(q, autoescape=True)
    with engine.connect() as conn:
        total = conn.scalar(select(func.count()).select_from(records).where(criteria))
        rows = conn.execute(select(records.c.id, records.c.nombre, records.c.rut, records.c.fecha).where(criteria).order_by(records.c.id).limit(25).offset((page-1)*25)).mappings().all()
        all_total = conn.scalar(select(func.count()).select_from(records))
    return render_template('admin.html', rows=rows, total=total, all_total=all_total, page=page, q=q)

@app.post('/admin/nomina/importar')
@admin
def import_roster():
    try:
        upload = request.files.get('archivo')
        if not upload or not upload.filename: raise ValueError('Selecciona un archivo CSV o XLSX.')
        if request.form.get('confirmar') != 'si': raise ValueError('Confirma que esta es la nómina completa autorizada.')
        raw = upload.read()
        if upload.filename.lower().endswith('.xlsx'):
            from openpyxl import load_workbook
            with zipfile.ZipFile(io.BytesIO(raw)) as z:
                if sum(i.file_size for i in z.infolist()) > 30_000_000: raise ValueError('El Excel supera el tamaño permitido.')
            wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            try:
                rows = list(wb.active.iter_rows(values_only=True))
            finally: wb.close()
        elif upload.filename.lower().endswith('.csv'):
            try: text = raw.decode('utf-8-sig')
            except UnicodeDecodeError: text = raw.decode('cp1252')
            try: dialect = csv.Sniffer().sniff(text[:8000], delimiters=',;\t')
            except csv.Error: dialect = csv.excel
            rows = list(csv.reader(io.StringIO(text), dialect))
        else: raise ValueError('Usa CSV o Excel XLSX (no XLS).')
        rows = [row for row in rows if any(v is not None and str(v).strip() for v in row)]
        if len(rows) < 2 or len(rows) > 20001: raise ValueError('La nómina debe tener encabezados y entre 1 y 20.000 personas.')
        headings = [' '.join(name_key(str(v or ''))) for v in rows[0]]
        rut_col = next((i for i,v in enumerate(headings) if v == 'RUT'), None)
        name_col = next((i for i,v in enumerate(headings) if v in ('NOMBRE', 'COMPLETO NOMBRE', 'NOMBRES')), None)
        if rut_col is None or name_col is None: raise ValueError('El archivo debe tener columnas RUT y NOMBRE (o Nombre completo). Correlativo es opcional.')
        incoming = {}; errors = []
        for n, row in enumerate(rows[1:], 2):
            try:
                if max(rut_col,name_col) >= len(row): raise ValueError('Faltan columnas.')
                rut = normalize_rut(str(row[rut_col] or ''))
                name = ' '.join(str(row[name_col] or '').split())
                if len(name) > 180 or len(name_key(name)) < 2: raise ValueError('Nombre completo inválido.')
                if rut in incoming: raise ValueError('RUT repetido dentro del archivo.')
                incoming[rut] = name
            except ValueError as e:
                errors.append(f'Fila {n}: {e}')
        if errors: raise ValueError('No se importó ningún dato. '+ ' / '.join(errors[:12]) + (f' / Total: {len(errors)} filas con errores.' if len(errors)>12 else ''))
        with engine.begin() as conn:
            existing = conn.scalar(select(func.count()).select_from(members))
            if existing and not hmac.compare_digest(request.form.get('password',''), os.environ['DELETE_PASSWORD']):
                raise ValueError('Para sustituir la nómina existente ingresa la clave de eliminación.')
            conn.execute(delete(members))
            conn.execute(members.insert(), [{'rut':r,'nombre':n} for r,n in incoming.items()])
        flash(f'Nómina importada: {len(incoming)} socios habilitados. Las firmas anteriores se conservaron.')
    except (ValueError, zipfile.BadZipFile, OSError, KeyError):
        import sys
        error = sys.exc_info()[1]
        flash(str(error) if isinstance(error, ValueError) else 'No se pudo leer el archivo. Usa un Excel XLSX o CSV válido.')
    return redirect(url_for('roster_report'))

def report_query(state, q=''):
    joined = members.outerjoin(records, members.c.rut == records.c.rut)
    query = select(members.c.nombre, members.c.rut, records.c.fecha).select_from(joined)
    if state == 'pendientes': query = query.where(records.c.id.is_(None))
    if state == 'registrados': query = query.where(records.c.id.is_not(None))
    if q: query = query.where(members.c.nombre.ilike('%'+q.replace('%','\\%').replace('_','\\_')+'%', escape='\\') | members.c.rut.contains(q, autoescape=True))
    return query.order_by(members.c.nombre, members.c.rut)

@app.get('/admin/reporte')
@admin
def roster_report():
    state = request.args.get('estado', 'pendientes')
    if state not in ('pendientes','registrados','todos'): abort(400)
    q = request.args.get('q','').strip()[:180]
    page = max(1, request.args.get('page', 1, type=int))
    with engine.connect() as conn:
        total_members = conn.scalar(select(func.count()).select_from(members))
        signed = conn.scalar(select(func.count()).select_from(members.join(records, members.c.rut == records.c.rut)))
        outside = conn.scalar(select(func.count()).select_from(records.outerjoin(members, records.c.rut == members.c.rut)).where(members.c.rut.is_(None)))
        query = report_query(state, q)
        total = conn.scalar(select(func.count()).select_from(query.subquery()))
        rows = conn.execute(query.limit(50).offset((page-1)*50)).mappings().all()
    return render_template('reporte.html', rows=rows, total_members=total_members, signed=signed,
                           pending=total_members-signed, outside=outside, total=total, page=page, state=state, q=q)

@app.get('/admin/reporte.csv')
@admin
def export_report():
    state = request.args.get('estado', 'pendientes')
    if state not in ('pendientes','registrados','todos'): abort(400)
    with engine.connect() as conn: rows = conn.execute(report_query(state, request.args.get('q','').strip()[:180])).mappings().all()
    output = io.StringIO(); writer = csv.writer(output, delimiter=';')
    writer.writerow(['Nombre completo','RUT','Estado','Fecha de registro (UTC)'])
    for item in rows:
        # Prevent spreadsheet formula interpretation of names.
        name = item['nombre']
        if name.startswith(('=','+','-','@')): name = "'" + name
        writer.writerow([name, item['rut'], 'Registrado' if item['fecha'] else 'Pendiente', item['fecha'].isoformat() if item['fecha'] else ''])
    return send_file(io.BytesIO(output.getvalue().encode('utf-8-sig')), as_attachment=True,
                     download_name='Reporte_'+state+'.csv', mimetype='text/csv; charset=utf-8')

@app.get('/admin/firma/<int:record_id>')
@admin
def get_signature(record_id):
    with engine.connect() as conn: data = conn.scalar(select(records.c.firma).where(records.c.id == record_id))
    if data is None: abort(404)
    return send_file(io.BytesIO(data), mimetype='image/png')

@app.post('/admin/eliminar/<int:record_id>')
@admin
def remove(record_id):
    if not hmac.compare_digest(request.form.get('password',''), os.environ['DELETE_PASSWORD']):
        flash('Clave de eliminación incorrecta.'); return redirect(url_for('dashboard'))
    with engine.begin() as conn: conn.execute(delete(records).where(records.c.id == record_id))
    flash('Registro eliminado.'); return redirect(url_for('dashboard'))

@app.post('/admin/logout')
def logout(): session.clear(); return redirect(url_for('index'))

def make_word(rows, date_text=None):
    doc = Document(ROOT / 'plantillas/censura.docx'); table = doc.tables[0]
    # Keep the annex heading with its introduction; remove trailing empty pages.
    for p in doc.paragraphs:
        if p.text == 'SOCIOS SOLICITANTES': p.paragraph_format.keep_with_next = True
    while doc.paragraphs and not doc.paragraphs[-1].text.strip():
        p = doc.paragraphs[-1]._p; p.getparent().remove(p)
    prototype = deepcopy(table.rows[1]._tr)
    for row in list(table.rows)[1:]: table._tbl.remove(row._tr)
    first_table = table
    for n, item in enumerate(rows, 1):
        if n > 1 and (n-1) % 10 == 0:
            separator = OxmlElement('w:p'); run = OxmlElement('w:r'); br = OxmlElement('w:br')
            from docx.oxml.ns import qn
            br.set(qn('w:type'), 'page'); run.append(br); separator.append(run)
            table._tbl.addnext(separator)
            new_tbl = deepcopy(first_table._tbl)
            for tr in list(new_tbl)[1:]:
                if tr.tag == qn('w:tr') and tr is not new_tbl.findall(qn('w:tr'))[0]: new_tbl.remove(tr)
            separator.addnext(new_tbl); table = WordTable(new_tbl, doc._body)
        table._tbl.append(deepcopy(prototype)); cells = table.rows[-1].cells
        for cell in cells: cell.text = ''
        cells[0].text = str(n); cells[1].text = item['nombre']; cells[2].text = item['rut']
        cells[3].paragraphs[0].add_run().add_picture(io.BytesIO(item['firma']), width=Inches(1.15), height=Inches(.42))
        trpr = table.rows[-1]._tr.get_or_add_trPr(); trpr.append(OxmlElement('w:cantSplit'))
    while len(table.rows) < 11:
        table._tbl.append(deepcopy(prototype))
        for cell in table.rows[-1].cells: cell.text = ''
    for t in doc.tables:
        header = t.rows[0]._tr.get_or_add_trPr(); header.append(OxmlElement('w:tblHeader'))
    for p in doc.paragraphs:
        if p.text == 'DATOS DE CONTACTO PARA LA COORDINACIÓN DEL PROCEDIMIENTO':
            p.paragraph_format.page_break_before = True
    if date_text:
        for p in doc.paragraphs:
            if 'En Antofagasta, a XX de octubre de 2026.' in p.text:
                for run in p.runs:
                    if 'XX' in run.text: run.text = run.text.replace('XX de octubre de 2026', date_text)
                # Handle templates with the date split into several runs.
                if 'XX' in p.text: p.text = 'En Antofagasta, a ' + date_text + '.'
    out = io.BytesIO(); doc.save(out); out.seek(0); return out

@app.get('/admin/word')
@admin
def export():
    date_text = None
    if request.args.get('fecha'):
        try:
            d = datetime.strptime(request.args['fecha'], '%Y-%m-%d'); months = ['enero','febrero','marzo','abril','mayo','junio','julio','agosto','septiembre','octubre','noviembre','diciembre']
            date_text = f'{d.day} de {months[d.month-1]} de {d.year}'
        except ValueError: abort(400)
    with engine.connect() as conn: rows = conn.execute(select(records).order_by(records.c.id)).mappings().all()
    return send_file(make_word(rows, date_text), as_attachment=True, download_name='Solicitud_censura_firmada.docx', mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document')

@app.get('/health')
def health():
    with engine.connect() as conn: conn.execute(select(1))
    return {'status':'ok'}

if __name__ == '__main__': app.run(port=int(os.getenv('PORT', 5000)))
